"""Reading what arrives: the email thread and its attachments.

Spreadsheets are read by code, never by the reasoning engine. A bill of
materials in a spreadsheet already has structure, and a model reading it can
only lose rows or move numbers between columns. The engine is used for the
free text, where structure genuinely has to be inferred.

Intake is repeat-safe. The same email, delivered twice, produces one case and
one set of lines, because the sender's own message identifier is unique in the
database.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from ..domain.units import canonical_unit
from ..money import dec
from ..text import count

# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------
#: Column headings seen in real bills of material, in English and Chinese.
HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "line_number": ("item", "item no", "no", "no.", "line", "line no", "pos", "seq", "項次", "序號"),
    "description": (
        "description", "part description", "specification", "spec", "material",
        "part", "part name", "item description", "product", "品名", "規格", "說明",
    ),
    "part_number": (
        "part number", "part no", "part no.", "p/n", "pn", "customer part",
        "your part", "mpn", "manufacturer part", "料號", "零件號",
    ),
    "quantity": ("qty", "quantity", "q'ty", "qty.", "order qty", "pcs", "數量"),
    "unit": ("unit", "uom", "u/m", "單位"),
    "requested_date": ("required date", "need by", "delivery date", "due date", "交期", "需求日"),
    "note": ("remark", "remarks", "note", "notes", "comment", "備註"),
}

_HEADER_LOOKUP = {
    synonym: field_name
    for field_name, synonyms in HEADER_SYNONYMS.items()
    for synonym in synonyms
}


@dataclass(slots=True)
class ExtractedRow:
    line_number: int
    description: str
    customer_reference: str
    quantity: Decimal
    unit: str
    requested_date: str
    note: str
    source_reference: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "line_number": self.line_number,
            "description": self.description,
            "customer_reference": self.customer_reference,
            "quantity": str(self.quantity),
            "unit": self.unit,
            "requested_date": self.requested_date,
            "note": self.note,
            "source_reference": self.source_reference,
        }


@dataclass(slots=True)
class AttachmentReading:
    filename: str
    parsed_with: str
    rows: list[ExtractedRow] = field(default_factory=list)
    text: str = ""
    note: str = ""
    problems: list[str] = field(default_factory=list)


def _normalise_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower()).rstrip(":：")


def _find_header_row(grid: list[list[Any]]) -> tuple[int, dict[int, str]] | None:
    """Find the row that looks like column headings, within the first 15 rows.

    Real bills of material start with a company logo, an address block and a
    couple of blank rows, so the headings are rarely in row one.
    """
    best: tuple[int, dict[int, str]] | None = None
    best_score = 0
    for index, row in enumerate(grid[:15]):
        mapping: dict[int, str] = {}
        for column, cell in enumerate(row):
            field_name = _HEADER_LOOKUP.get(_normalise_header(cell))
            if field_name and field_name not in mapping.values():
                mapping[column] = field_name
        score = len(mapping) + (2 if "quantity" in mapping.values() else 0)
        if "description" in mapping.values() and score > best_score:
            best, best_score = (index, mapping), score
    return best


def _parse_quantity(value: Any) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        return dec(value)
    text = re.sub(r"[^\d.\-]", "", str(value))
    return dec(text or 0)


def _parse_date(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    match = re.search(r"(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})", str(value))
    if match:
        year, month, day = match.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"
    return ""


def rows_from_grid(grid: list[list[Any]], *, sheet_name: str) -> AttachmentReading:
    reading = AttachmentReading(filename=sheet_name, parsed_with="spreadsheet_reader")
    found = _find_header_row(grid)
    if found is None:
        reading.problems.append(
            "No column headings were recognised, so the sheet was not read as a "
            "bill of materials."
        )
        return reading

    header_index, mapping = found
    reading.note = (
        "Headings found on row "
        + str(header_index + 1)
        + ": "
        + ", ".join(sorted(set(mapping.values())))
    )

    line_number = 0
    for offset, row in enumerate(grid[header_index + 1 :], start=header_index + 2):
        values = {
            field_name: (row[column] if column < len(row) else None)
            for column, field_name in mapping.items()
        }
        description = str(values.get("description") or "").strip()
        part_number = str(values.get("part_number") or "").strip()
        if not description and not part_number:
            continue
        quantity = _parse_quantity(values.get("quantity"))
        if not description:
            description = part_number
        elif part_number and part_number.lower() not in description.lower():
            description = f"{part_number} {description}"

        line_number += 1
        reading.rows.append(
            ExtractedRow(
                line_number=line_number,
                description=re.sub(r"\s+", " ", description),
                customer_reference=part_number,
                quantity=quantity,
                unit=canonical_unit(str(values.get("unit") or "piece")),
                requested_date=_parse_date(values.get("requested_date")),
                note=str(values.get("note") or "").strip(),
                source_reference=f"{sheet_name}, row {offset}",
            )
        )
        if quantity <= 0:
            reading.problems.append(
                f"{sheet_name} row {offset} has no quantity, so the line is flagged."
            )

    if not reading.rows:
        reading.problems.append("The headings were found but no rows had a description.")
    return reading


def read_spreadsheet(data: bytes, filename: str) -> AttachmentReading:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - depends on the install
        return AttachmentReading(
            filename=filename,
            parsed_with="unavailable",
            problems=[
                "The spreadsheet reader is not installed. Run "
                "pip install -r backend/requirements.txt."
            ],
        )

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    combined = AttachmentReading(filename=filename, parsed_with="spreadsheet_reader")
    for sheet in workbook.worksheets:
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        if not grid:
            continue
        reading = rows_from_grid(grid, sheet_name=f"{filename}, sheet {sheet.title}")
        offset = len(combined.rows)
        for row in reading.rows:
            row.line_number += offset
        combined.rows.extend(reading.rows)
        combined.problems.extend(reading.problems)
        if reading.note:
            combined.note = (combined.note + " " + reading.note).strip()
    workbook.close()
    if not combined.rows and not combined.problems:
        combined.problems.append("The workbook has no readable rows.")
    return combined


def read_csv(data: bytes, filename: str) -> AttachmentReading:
    text = data.decode("utf-8-sig", errors="replace")
    dialect: Any = csv.excel
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        pass
    grid = [row for row in csv.reader(io.StringIO(text), dialect)]
    return rows_from_grid(grid, sheet_name=filename)


def read_pdf(data: bytes, filename: str) -> AttachmentReading:
    """Pull the text layer out of a PDF. Scanned paper is reported, not guessed."""
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - depends on the install
        return AttachmentReading(
            filename=filename,
            parsed_with="unavailable",
            problems=["The PDF reader is not installed."],
        )

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    reading = AttachmentReading(
        filename=filename,
        parsed_with="pdf_text_reader",
        text=text,
        note=f"{count(len(reader.pages), 'page')} read.",
    )
    if not text:
        reading.problems.append(
            "This PDF has no text layer, so it is probably a scan. Reading scanned "
            "paper is outside the scope of this version; the line is flagged for a "
            "person to read."
        )
    return reading


def read_attachment(data: bytes, filename: str, content_type: str = "") -> AttachmentReading:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return read_spreadsheet(data, filename)
    if lower.endswith(".csv") or lower.endswith(".tsv"):
        return read_csv(data, filename)
    if lower.endswith(".pdf"):
        return read_pdf(data, filename)
    if lower.endswith((".txt", ".md")):
        return AttachmentReading(
            filename=filename,
            parsed_with="plain_text",
            text=data.decode("utf-8", errors="replace"),
        )
    if lower.endswith(".xls"):
        return AttachmentReading(
            filename=filename,
            parsed_with="unavailable",
            problems=[
                "This is the old Excel format. Ask the customer for an .xlsx file, "
                "or convert it before loading."
            ],
        )
    return AttachmentReading(
        filename=filename,
        parsed_with="unavailable",
        problems=[f"{filename} is a kind of file this version does not read."],
    )


# ---------------------------------------------------------------------------
# Email threads
# ---------------------------------------------------------------------------
_QUOTE_MARKER = re.compile(
    r"^(>|On .+ wrote:|-{2,}\s*Original Message|From:\s|寄件者:|寄件人:)", re.M
)


def split_thread(body: str) -> tuple[str, str]:
    """Separate the newest message from the quoted history below it.

    The whole thread still gets read, because the buyer's colleague often adds
    the important detail three replies down, but the newest part is what the
    reply is written against.
    """
    match = _QUOTE_MARKER.search(body or "")
    if not match:
        return (body or "").strip(), ""
    return body[: match.start()].strip(), body[match.start() :].strip()


_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")


def addresses_in(text: str) -> list[str]:
    seen: list[str] = []
    for address in _EMAIL_RE.findall(text or ""):
        lowered = address.lower()
        if lowered not in seen:
            seen.append(lowered)
    return seen
